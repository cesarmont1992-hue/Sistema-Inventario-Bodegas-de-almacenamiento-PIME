import tkinter as tk
from tkinter import messagebox, filedialog
import os
from datetime import datetime

from PIL import Image, ImageTk
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.units import mm

from inventario import InventarioVentana
from configuracion import ConfiguracionVentana
from database.conexion import conectar_bd, crear_tablas

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

ventana = tk.Tk()
ventana.title("Sistema de Inventario")
ventana.geometry("1000x650")
ventana.minsize(900, 600)
ventana.configure(bg="#F5F1EA")

ICONO = os.path.join(BASE_DIR, "assets", "warehouse.ico")
if os.path.exists(ICONO):
    try:
        ventana.iconbitmap(ICONO)
    except Exception as error:
        print("No se pudo cargar el icono:", error)

ventana_inventario = None
ventana_configuracion = None
panel_notificaciones = None


def centrar_ventana():
    ventana.update_idletasks()
    ancho, alto = 1000, 650
    x = (ventana.winfo_screenwidth() - ancho) // 2
    y = (ventana.winfo_screenheight() - alto) // 2
    ventana.geometry(f"{ancho}x{alto}+{x}+{y}")


def asegurar_tabla_configuracion():
    conexion = None
    try:
        conexion = conectar_bd()
        cursor = conexion.cursor()
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS configuracion (
                id INTEGER PRIMARY KEY,
                nombre_empresa TEXT DEFAULT '',
                nit TEXT DEFAULT '',
                telefono TEXT DEFAULT '',
                direccion TEXT DEFAULT '',
                iva REAL DEFAULT 19,
                moneda TEXT DEFAULT 'COP',
                prefijo_factura TEXT DEFAULT 'FAC',
                numero_factura INTEGER DEFAULT 1,
                stock_minimo_alerta REAL DEFAULT 5
            )
        """)
        cursor.execute("SELECT id FROM configuracion WHERE id = 1")
        if cursor.fetchone() is None:
            cursor.execute("""
                INSERT INTO configuracion
                (id, nombre_empresa, nit, telefono, direccion, iva,
                 moneda, prefijo_factura, numero_factura, stock_minimo_alerta)
                VALUES (1, '', '', '', '', 19, 'COP', 'FAC', 1, 5)
            """)
        conexion.commit()
    except Exception as error:
        if conexion:
            conexion.rollback()
        messagebox.showerror("Error de configuración", str(error), parent=ventana)
    finally:
        if conexion:
            conexion.close()


def obtener_configuracion():
    conexion = None
    try:
        conexion = conectar_bd()
        cursor = conexion.cursor()
        cursor.execute("""
            SELECT nombre_empresa, nit, telefono, direccion, iva, moneda,
                   prefijo_factura, numero_factura, stock_minimo_alerta
            FROM configuracion WHERE id = 1
        """)
        return cursor.fetchone()
    except Exception as error:
        print("Error al cargar configuración:", error)
        return None
    finally:
        if conexion:
            conexion.close()


def actualizar_datos_empresa():
    datos = obtener_configuracion()
    if datos:
        etiqueta_empresa.config(text=datos[0] or "Empresa sin configurar")
        etiqueta_datos.config(
            text=(
                f"NIT: {datos[1] or 'No registrado'}     |     "
                f"Teléfono: {datos[2] or 'No registrado'}\n"
                f"Dirección: {datos[3] or 'No registrada'}"
            )
        )
    else:
        etiqueta_empresa.config(text="Empresa sin configurar")
        etiqueta_datos.config(text="Ingresa los datos de la empresa desde Configuración")


def obtener_datos_reporte():
    conexion = None
    try:
        conexion = conectar_bd()
        cursor = conexion.cursor()
        cursor.execute("SELECT COUNT(*) FROM productos WHERE activo = 1")
        cantidad_productos = cursor.fetchone()[0] or 0
        cursor.execute("""
            SELECT COALESCE(SUM(i.stock_actual), 0)
            FROM inventario i
            INNER JOIN productos p ON p.id_producto = i.id_producto
            WHERE p.activo = 1
        """)
        total_unidades = cursor.fetchone()[0] or 0
        cursor.execute("SELECT COUNT(*) FROM marcas WHERE activo = 1")
        cantidad_marcas = cursor.fetchone()[0] or 0
        cursor.execute("""
            SELECT COALESCE(SUM(i.stock_actual * p.precio_compra), 0)
            FROM inventario i
            INNER JOIN productos p ON p.id_producto = i.id_producto
            WHERE p.activo = 1
        """)
        valor_total = cursor.fetchone()[0] or 0
        cursor.execute("""
            SELECT p.codigo, p.codigo_barras, p.nombre,
                   COALESCE(m.nombre, ''), COALESCE(i.stock_actual, 0),
                   COALESCE(p.precio_compra, 0),
                   COALESCE(i.stock_actual, 0) * COALESCE(p.precio_compra, 0)
            FROM productos p
            LEFT JOIN marcas m ON p.id_marca = m.id_marca
            LEFT JOIN inventario i ON p.id_producto = i.id_producto
            WHERE p.activo = 1 ORDER BY p.nombre
        """)
        return cantidad_productos, total_unidades, cantidad_marcas, valor_total, cursor.fetchall()
    except Exception as error:
        messagebox.showerror("Error", f"No fue posible obtener los datos.\n\n{error}", parent=ventana)
        return None
    finally:
        if conexion:
            conexion.close()


def formato_dinero(valor):
    try:
        return f"$ {valor:,.0f}"
    except (TypeError, ValueError):
        return "$ 0"


def exportar_excel():
    datos = obtener_datos_reporte()
    if datos is None:
        return
    cantidad_productos, total_unidades, cantidad_marcas, valor_total, detalle = datos
    configuracion = obtener_configuracion()
    empresa = configuracion[0] if configuracion else "Empresa sin configurar"
    nit = configuracion[1] if configuracion else ""
    fecha = datetime.now().strftime("%Y-%m-%d_%H-%M")
    ruta = filedialog.asksaveasfilename(
        parent=ventana, title="Guardar reporte Excel", defaultextension=".xlsx",
        initialfile=f"Reporte_Inventario_{fecha}.xlsx",
        filetypes=[("Archivo Excel", "*.xlsx")]
    )
    if not ruta:
        return
    try:
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Resumen"
        hoja.merge_cells("A1:B1")
        hoja["A1"] = "REPORTE DE INVENTARIO"
        hoja["A1"].font = Font(bold=True, size=18, color="FFFFFF")
        hoja["A1"].fill = PatternFill("solid", fgColor="241F1A")
        hoja["A1"].alignment = Alignment(horizontal="center")
        hoja["A3"], hoja["B3"] = "Empresa", empresa
        hoja["A4"], hoja["B4"] = "NIT", nit
        hoja["A5"], hoja["B5"] = "Fecha", datetime.now().strftime("%d/%m/%Y %H:%M")
        hoja["A7"], hoja["B7"] = "Indicador", "Resultado"
        for celda in hoja[7]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="8A5A2B")
            celda.alignment = Alignment(horizontal="center")
        for fila, (indicador, resultado) in enumerate([
            ("Productos registrados", cantidad_productos),
            ("Unidades en inventario", total_unidades),
            ("Marcas registradas", cantidad_marcas),
            ("Valor total de mercancía", formato_dinero(valor_total)),
        ], start=8):
            hoja.cell(fila, 1, indicador)
            hoja.cell(fila, 2, resultado)
        hoja.column_dimensions["A"].width = 32
        hoja.column_dimensions["B"].width = 30

        detalle_hoja = libro.create_sheet("Detalle Inventario")
        encabezados = ["Código", "Código de barras", "Producto", "Marca", "Stock", "Precio compra", "Valor total"]
        for columna, encabezado in enumerate(encabezados, 1):
            celda = detalle_hoja.cell(1, columna, encabezado)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="8A5A2B")
            celda.alignment = Alignment(horizontal="center")
        for fila, producto in enumerate(detalle, 2):
            for columna, valor in enumerate(producto, 1):
                detalle_hoja.cell(fila, columna, valor)
            detalle_hoja.cell(fila, 6).number_format = "$ #,##0"
            detalle_hoja.cell(fila, 7).number_format = "$ #,##0"
        for columna, ancho in {1:15, 2:20, 3:35, 4:25, 5:12, 6:18, 7:20}.items():
            detalle_hoja.column_dimensions[get_column_letter(columna)].width = ancho
        detalle_hoja.freeze_panes = "A2"
        libro.save(ruta)
        messagebox.showinfo("Excel generado", f"El reporte Excel fue generado correctamente.\n\nArchivo:\n{ruta}", parent=ventana)
    except Exception as error:
        messagebox.showerror("Error", f"No fue posible generar el Excel.\n\n{error}", parent=ventana)


def generar_pdf():
    datos = obtener_datos_reporte()
    if datos is None:
        return
    cantidad_productos, total_unidades, cantidad_marcas, valor_total, detalle = datos
    configuracion = obtener_configuracion()
    empresa = configuracion[0] if configuracion else "Empresa sin configurar"
    nit = configuracion[1] if configuracion else ""
    telefono = configuracion[2] if configuracion else ""
    direccion = configuracion[3] if configuracion else ""
    fecha = datetime.now().strftime("%Y-%m-%d_%H-%M")
    ruta = filedialog.asksaveasfilename(
        parent=ventana, title="Guardar reporte PDF", defaultextension=".pdf",
        initialfile=f"Reporte_Inventario_{fecha}.pdf",
        filetypes=[("Archivo PDF", "*.pdf")]
    )
    if not ruta:
        return
    try:
        documento = SimpleDocTemplate(ruta, pagesize=landscape(A4), rightMargin=12*mm, leftMargin=12*mm, topMargin=12*mm, bottomMargin=12*mm)
        estilos = getSampleStyleSheet()
        titulo = ParagraphStyle("Titulo", parent=estilos["Title"], fontSize=20, alignment=TA_CENTER, textColor=colors.HexColor("#241F1A"), spaceAfter=8)
        empresa_style = ParagraphStyle("Empresa", parent=estilos["Normal"], fontSize=14, alignment=TA_CENTER, textColor=colors.HexColor("#8A5A2B"), spaceAfter=5)
        normal = ParagraphStyle("NormalPersonalizado", parent=estilos["Normal"], fontSize=9, leading=12)
        contenido = [Paragraph("REPORTE DE INVENTARIO", titulo), Paragraph(empresa, empresa_style)]
        contenido.append(Paragraph(f"NIT: {nit} &nbsp;&nbsp;&nbsp; Teléfono: {telefono} &nbsp;&nbsp;&nbsp; Dirección: {direccion}", normal))
        contenido.append(Spacer(1, 8))
        contenido.append(Paragraph("Fecha del reporte: " + datetime.now().strftime("%d/%m/%Y %H:%M"), normal))
        contenido.append(Spacer(1, 15))
        resumen = Table([
            ["Productos", "Unidades", "Marcas", "Valor total mercancía"],
            [str(cantidad_productos), f"{total_unidades:,.0f}", str(cantidad_marcas), formato_dinero(valor_total)]
        ], colWidths=[55*mm, 55*mm, 55*mm, 70*mm])
        resumen.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#8A5A2B")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.white),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("ALIGN", (0,0), (-1,-1), "CENTER"),
            ("BACKGROUND", (0,1), (-1,1), colors.HexColor("#F5F1EA")),
            ("BOX", (0,0), (-1,-1), 0.8, colors.HexColor("#8A5A2B")),
            ("INNERGRID", (0,0), (-1,-1), 0.4, colors.grey),
            ("TOPPADDING", (0,0), (-1,-1), 8), ("BOTTOMPADDING", (0,0), (-1,-1), 8)
        ]))
        contenido += [resumen, Spacer(1, 18), Paragraph("DETALLE DEL INVENTARIO", estilos["Heading2"])]
        tabla = [["Código", "Código barras", "Producto", "Marca", "Stock", "Precio compra", "Valor total"]]
        for codigo, barras, nombre, marca, stock, compra, total in detalle:
            tabla.append([str(codigo or ""), str(barras or ""), str(nombre or ""), str(marca or ""), f"{stock:,.0f}", formato_dinero(compra), formato_dinero(total)])
        if len(tabla) == 1:
            tabla.append(["", "", "No hay productos registrados", "", "", "", ""])
        detalle_tabla = Table(tabla, repeatRows=1, colWidths=[25*mm, 32*mm, 55*mm, 38*mm, 20*mm, 35*mm, 38*mm])
        detalle_tabla.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#241F1A")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.white),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE", (0,0), (-1,-1), 7),
            ("ALIGN", (0,0), (-1,-1), "CENTER"),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("GRID", (0,0), (-1,-1), 0.4, colors.grey),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F5F1EA")]),
            ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5)
        ]))
        contenido += [detalle_tabla, Spacer(1, 12), Paragraph("Reporte generado por Sistema de Inventario", ParagraphStyle("Pie", parent=estilos["Normal"], fontSize=8, alignment=TA_CENTER, textColor=colors.grey))]
        documento.build(contenido)
        messagebox.showinfo("PDF generado", f"El reporte PDF fue generado correctamente.\n\nArchivo:\n{ruta}", parent=ventana)
    except Exception as error:
        messagebox.showerror("Error", f"No fue posible generar el PDF.\n\n{error}", parent=ventana)


def abrir_inventario():
    global ventana_inventario
    if ventana_inventario is not None:
        try:
            if ventana_inventario.winfo_exists():
                ventana_inventario.lift(); ventana_inventario.focus_force(); return
        except tk.TclError:
            ventana_inventario = None
    ventana_inventario = InventarioVentana(ventana)
    def cerrar():
        global ventana_inventario
        try: ventana_inventario.destroy()
        except tk.TclError: pass
        ventana_inventario = None
    ventana_inventario.protocol("WM_DELETE_WINDOW", cerrar)


def abrir_configuracion():
    global ventana_configuracion
    if ventana_configuracion is not None:
        try:
            if ventana_configuracion.winfo_exists():
                ventana_configuracion.lift(); ventana_configuracion.focus_force(); return
        except tk.TclError:
            ventana_configuracion = None
    ventana_configuracion = ConfiguracionVentana(ventana)
    def cerrar():
        global ventana_configuracion
        try: ventana_configuracion.destroy()
        except tk.TclError: pass
        ventana_configuracion = None
        actualizar_datos_empresa()
    ventana_configuracion.protocol("WM_DELETE_WINDOW", cerrar)



def obtener_alertas_principales():
    """Consulta directamente las alertas de stock desde SQLite."""
    conexion = None
    try:
        conexion = conectar_bd()
        cursor = conexion.cursor()
        cursor.execute("""
            SELECT
                p.id_producto,
                p.nombre,
                COALESCE(i.stock_actual, 0),
                COALESCE(i.stock_minimo, 0),
                COALESCE(i.stock_maximo, 0)
            FROM productos p
            INNER JOIN inventario i ON i.id_producto = p.id_producto
            WHERE p.activo = 1
              AND (
                    COALESCE(i.stock_actual, 0) <= 0
                    OR COALESCE(i.stock_actual, 0) < COALESCE(i.stock_minimo, 0)
                    OR (
                        COALESCE(i.stock_maximo, 0) > 0
                        AND COALESCE(i.stock_actual, 0) > COALESCE(i.stock_maximo, 0)
                    )
              )
            ORDER BY
                CASE
                    WHEN COALESCE(i.stock_actual, 0) <= 0 THEN 1
                    WHEN COALESCE(i.stock_actual, 0) < COALESCE(i.stock_minimo, 0) THEN 2
                    ELSE 3
                END,
                p.nombre
        """)
        resultado = []
        for id_producto, nombre, stock, minimo, maximo in cursor.fetchall():
            if stock <= 0:
                tipo = "AGOTADO"
                icono = "🔴"
            elif stock < minimo:
                tipo = "STOCK BAJO"
                icono = "🟠"
            else:
                tipo = "STOCK EXCEDIDO"
                icono = "🟡"
            resultado.append((id_producto, nombre, stock, minimo, maximo, tipo, icono))
        return resultado
    except Exception as error:
        print("Error al consultar alertas:", error)
        return []
    finally:
        if conexion:
            conexion.close()


def actualizar_campana_alertas():
    """Actualiza el contador de la campana sin abrir ventanas."""
    try:
        if not ventana.winfo_exists():
            return
    except tk.TclError:
        return

    alertas = obtener_alertas_principales()
    cantidad = len(alertas)

    if cantidad > 0:
        boton_campana.config(text="🔔")
        etiqueta_contador.config(text=str(cantidad), bg="#D93025", fg="white")
        etiqueta_contador.place(relx=1.0, x=-8, y=2, anchor="ne")
    else:
        boton_campana.config(text="🔔")
        etiqueta_contador.config(text="", bg="#D93025")
        etiqueta_contador.place_forget()

    ventana.after(10000, actualizar_campana_alertas)


def mostrar_notificaciones():
    """Muestra un panel de notificaciones estilo Outlook junto a la campana."""
    global panel_notificaciones

    if panel_notificaciones is not None:
        try:
            if panel_notificaciones.winfo_exists():
                panel_notificaciones.destroy()
                panel_notificaciones = None
                return
        except tk.TclError:
            panel_notificaciones = None

    alertas = obtener_alertas_principales()

    panel_notificaciones = tk.Toplevel(ventana)
    panel_notificaciones.title("Notificaciones")
    panel_notificaciones.overrideredirect(True)
    panel_notificaciones.configure(bg="#FFFFFF")
    panel_notificaciones.geometry("390x500")

    # Posicionar el panel debajo de la campana.
    ventana.update_idletasks()
    x = boton_campana.winfo_rootx() - 345
    y = boton_campana.winfo_rooty() + boton_campana.winfo_height() + 6
    panel_notificaciones.geometry(f"390x500+{max(5, x)}+{max(5, y)}")

    borde = tk.Frame(panel_notificaciones, bg="#D9D9D9", bd=1, relief="solid")
    borde.pack(fill="both", expand=True)

    cabecera = tk.Frame(borde, bg="#FFFFFF")
    cabecera.pack(fill="x", padx=18, pady=(15, 10))

    tk.Label(
        cabecera,
        text="Notificaciones",
        bg="#FFFFFF",
        fg="#242424",
        font=("Arial", 16, "bold")
    ).pack(side="left")

    tk.Button(
        cabecera,
        text="✕",
        command=lambda: cerrar_notificaciones(),
        bg="#FFFFFF",
        fg="#666666",
        activebackground="#F2F2F2",
        activeforeground="#222222",
        relief="flat",
        bd=0,
        font=("Arial", 12),
        cursor="hand2"
    ).pack(side="right")

    tk.Frame(borde, bg="#E6E6E6", height=1).pack(fill="x")

    contenido_notificaciones = tk.Frame(borde, bg="#F8F8F8")
    contenido_notificaciones.pack(fill="both", expand=True)

    if not alertas:
        tk.Label(
            contenido_notificaciones,
            text="🔔",
            bg="#F8F8F8",
            fg="#777777",
            font=("Arial", 32)
        ).pack(pady=(85, 8))
        tk.Label(
            contenido_notificaciones,
            text="No tienes notificaciones",
            bg="#F8F8F8",
            fg="#333333",
            font=("Arial", 12, "bold")
        ).pack()
        tk.Label(
            contenido_notificaciones,
            text="El inventario no presenta alertas.",
            bg="#F8F8F8",
            fg="#777777",
            font=("Arial", 10)
        ).pack(pady=5)
    else:
        tk.Label(
            contenido_notificaciones,
            text=f"{len(alertas)} alerta(s) requieren tu atención",
            bg="#F8F8F8",
            fg="#555555",
            font=("Arial", 10)
        ).pack(anchor="w", padx=18, pady=(12, 7))

        canvas = tk.Canvas(contenido_notificaciones, bg="#F8F8F8", highlightthickness=0)
        scroll = ttk.Scrollbar(contenido_notificaciones, orient="vertical", command=canvas.yview)
        lista = tk.Frame(canvas, bg="#F8F8F8")
        canvas_window = canvas.create_window((0, 0), window=lista, anchor="nw", width=360)
        canvas.configure(yscrollcommand=scroll.set)

        def ajustar_scroll(event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))

        lista.bind("<Configure>", ajustar_scroll)
        canvas.bind("<Configure>", lambda e: canvas.itemconfigure(canvas_window, width=e.width))
        canvas.pack(side="left", fill="both", expand=True, padx=(12, 0), pady=(0, 10))
        scroll.pack(side="right", fill="y", padx=(0, 8), pady=(0, 10))

        for id_producto, nombre, stock, minimo, maximo, tipo, icono in alertas:
            tarjeta = tk.Frame(lista, bg="#FFFFFF", bd=1, relief="solid", cursor="hand2")
            tarjeta.pack(fill="x", padx=5, pady=4)

            fila = tk.Frame(tarjeta, bg="#FFFFFF", cursor="hand2")
            fila.pack(fill="x", padx=10, pady=8)

            tk.Label(
                fila, text=icono, bg="#FFFFFF", font=("Arial", 16), cursor="hand2"
            ).pack(side="left", padx=(0, 8))

            texto = tk.Frame(fila, bg="#FFFFFF", cursor="hand2")
            texto.pack(side="left", fill="x", expand=True)

            tk.Label(
                texto, text=tipo, bg="#FFFFFF", fg="#333333",
                font=("Arial", 9, "bold"), anchor="w", cursor="hand2"
            ).pack(fill="x")
            tk.Label(
                texto, text=nombre, bg="#FFFFFF", fg="#222222",
                font=("Arial", 10, "bold"), anchor="w", cursor="hand2"
            ).pack(fill="x")
            tk.Label(
                texto, text=f"Stock: {stock:g}   |   Mín.: {minimo:g}   |   Máx.: {maximo:g}",
                bg="#FFFFFF", fg="#777777", font=("Arial", 9), anchor="w", cursor="hand2"
            ).pack(fill="x", pady=(2, 0))

            def abrir_producto(event=None, pid=id_producto):
                cerrar_notificaciones()
                abrir_inventario()

            for widget in (tarjeta, fila, texto):
                widget.bind("<Button-1>", abrir_producto)

    pie = tk.Frame(borde, bg="#FFFFFF")
    pie.pack(fill="x")

    tk.Button(
        pie,
        text="Ver inventario",
        command=lambda: (cerrar_notificaciones(), abrir_inventario()),
        bg="#5B6F8C",
        fg="white",
        activebackground="#4A5C75",
        activeforeground="white",
        relief="flat",
        font=("Arial", 10, "bold"),
        cursor="hand2",
        padx=14,
        pady=8
    ).pack(pady=10)

    def cerrar_al_hacer_click(event):
        try:
            if event.widget == panel_notificaciones:
                cerrar_notificaciones()
        except tk.TclError:
            pass


def cerrar_notificaciones():
    global panel_notificaciones
    if panel_notificaciones is not None:
        try:
            panel_notificaciones.destroy()
        except tk.TclError:
            pass
        panel_notificaciones = None


def salir():
    if messagebox.askyesno("Salir", "¿Deseas salir del sistema?", parent=ventana):
        ventana.destroy()


# Menú lateral
menu = tk.Frame(ventana, bg="#241F1A", width=230)
menu.pack(side="left", fill="y")
menu.pack_propagate(False)

tk.Label(menu, text="INVENTARIO", bg="#241F1A", fg="white", font=("Arial", 20, "bold")).pack(pady=(40,10))
tk.Label(menu, text="Sistema de Gestión", bg="#241F1A", fg="#D8C7AE", font=("Arial", 10)).pack(pady=(0,30))

for texto, comando, fondo, pady in [
    ("📦  Inventario", abrir_inventario, "#3A3028", 5),
    ("⚙️  Configuración", abrir_configuracion, "#3A3028", 5),
    ("📊  Exportar Excel", exportar_excel, "#477A5B", 20),
    ("📄  Generar PDF", generar_pdf, "#8A5A2B", 5),
]:
    tk.Button(menu, text=texto, command=comando, bg=fondo, fg="white", activebackground=fondo, activeforeground="white", font=("Arial",11,"bold"), relief="flat", cursor="hand2", height=2).pack(fill="x", padx=15, pady=(pady,5))

tk.Button(menu, text="🚪  Salir", command=salir, bg="#8C3B3B", fg="white", activebackground="#A94A4A", activeforeground="white", font=("Arial",11,"bold"), relief="flat", cursor="hand2", height=2).pack(side="bottom", fill="x", padx=15, pady=20)

# Contenido principal
contenido = tk.Frame(ventana, bg="#F5F1EA")
contenido.pack(side="left", fill="both", expand=True)

# Barra superior con campana de notificaciones estilo Outlook
barra_superior = tk.Frame(contenido, bg="#F5F1EA", height=65)
barra_superior.pack(fill="x", padx=25, pady=(10, 0))
barra_superior.pack_propagate(False)

marco_campana = tk.Frame(barra_superior, bg="#F5F1EA", width=58, height=50)
marco_campana.pack(side="right", padx=(0, 5), pady=5)
marco_campana.pack_propagate(False)

boton_campana = tk.Button(
    marco_campana, text="🔔", command=mostrar_notificaciones,
    bg="#FFFFFF", fg="#333333", activebackground="#F2F2F2",
    activeforeground="#111111", relief="flat", bd=1,
    font=("Arial", 20), cursor="hand2"
)
boton_campana.pack(fill="both", expand=True)

etiqueta_contador = tk.Label(
    marco_campana, text="", bg="#D93025", fg="white",
    font=("Arial", 8, "bold"), padx=4, pady=1
)
etiqueta_contador.place_forget()

tk.Label(
    barra_superior, text="Sistema de Inventario", bg="#F5F1EA",
    fg="#241F1A", font=("Arial", 28, "bold")
).pack(side="left", pady=(8, 0))
etiqueta_empresa = tk.Label(contenido, text="Cargando empresa...", bg="#F5F1EA", fg="#8A5A2B", font=("Arial",19,"bold"))
etiqueta_empresa.pack(pady=(5,5))
etiqueta_datos = tk.Label(contenido, text="Cargando información...", bg="#F5F1EA", fg="#555555", font=("Arial",10))
etiqueta_datos.pack(pady=(0,5))
tk.Label(contenido, text="Control de productos y existencias", bg="#F5F1EA", fg="#555555", font=("Arial",13)).pack(pady=(5,0))

ruta_imagen = os.path.join(BASE_DIR, "assets", "bodega.jpg")
try:
    imagen = Image.open(ruta_imagen)
    imagen.thumbnail((500,320))
    imagen_tk = ImageTk.PhotoImage(imagen)
    etiqueta_imagen = tk.Label(contenido, image=imagen_tk, bg="#F5F1EA")
    etiqueta_imagen.image = imagen_tk
    etiqueta_imagen.pack(pady=25)
except Exception as error:
    tk.Label(contenido, text="No se pudo cargar la imagen\n\nVerifica que exista:\nassets/bodega.jpg", bg="#F5F1EA", fg="#8C3B3B", font=("Arial",12,"bold")).pack(pady=50)
    print("Error al cargar la imagen:", error)

try:
    crear_tablas()
    asegurar_tabla_configuracion()
    actualizar_datos_empresa()
except Exception as error:
    messagebox.showerror("Error de inicio", f"No fue posible preparar la base de datos.\n\n{error}", parent=ventana)

centrar_ventana()
ventana.after(500, actualizar_campana_alertas)
ventana.mainloop()
